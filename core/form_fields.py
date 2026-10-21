"""Criação de campos de formulário (AcroForm) em arquivos PDF.

Os campos são desenhados com o reportlab em um PDF de uma página (o que gera
os widgets AcroForm corretamente), mesclados na página alvo, e registrados no
dicionário /AcroForm do documento para que o campo fique realmente preenchível
em qualquer leitor de PDF.
"""

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import cast

from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    BooleanObject,
    DictionaryObject,
    IndirectObject,
    NameObject,
    StreamObject,
    read_object,
)
from reportlab.pdfgen import canvas

from .base import READ, PDFOperation, atomic_output, open_reader, require_valid_page_index, write_pdf
from .document import remove_widgets_from_form
from .messages import Message
from .tabular import escape_formula, write_table

_VALID_FIELD_TYPES = ("text", "checkbox", "date", "dropdown", "listbox")
_CHOICE_FIELD_TYPES = ("dropdown", "listbox")


class AddFormField(PDFOperation):
    """Adiciona um campo de formulário interativo (texto, data, caixa de seleção, lista
    suspensa ou lista de seleção) a uma página do PDF. Campo de data é um campo de texto
    comum com uma dica de formato no tooltip — PDF não tem um tipo de campo "data" nativo,
    e validar o formato exigiria uma ação de JavaScript, que este projeto evita adicionar
    de propósito (o mesmo motivo por trás de `core/pdfa.py` e `core/redaction.py` removerem
    JavaScript de PDFs de entrada)."""

    def run(
        self,
        input_path: str,
        output_path: str,
        field_name: str,
        page_number: int,
        x: float,
        y: float,
        width: float = 200,
        height: float = 20,
        field_type: str = "text",
        checked: bool = False,
        options: list[str] | None = None,
    ) -> None:
        if field_type not in _VALID_FIELD_TYPES:
            raise ValueError(f"field_type inválido: {field_type}")
        if field_type in _CHOICE_FIELD_TYPES and not options:
            raise ValueError(f"Informe ao menos uma opção (options) para o tipo {field_type}.")
        options = list(options or [])

        reader = open_reader(input_path)
        require_valid_page_index(len(reader.pages), page_number)
        _require_new_field_name(reader, field_name)

        writer = PdfWriter(clone_from=reader)
        target_page = writer.pages[page_number]
        page_width = float(target_page.mediabox.width)
        page_height = float(target_page.mediabox.height)

        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=(page_width, page_height))
        if field_type == "checkbox":
            c.acroForm.checkbox(name=field_name, x=x, y=y, size=height, checked=checked)
        elif field_type == "dropdown":
            # reportlab tem um bug conhecido (confirmado nesta sessão): choice()/listbox()
            # sem um `value` não-vazio deixam uma variável local (`lbextras`) sem atribuir
            # e derrubam com UnboundLocalError antes mesmo de desenhar o campo. Pré-selecionar
            # a primeira opção evita esse caminho de código, sem mudar o que o usuário vê
            # como padrão razoável (a maioria dos leitores já mostra a primeira opção).
            c.acroForm.choice(
                name=field_name,
                x=x,
                y=y,
                width=width,
                height=height,
                options=list(options),
                value=str(options[0]),
            )
        elif field_type == "listbox":
            c.acroForm.listbox(
                name=field_name,
                x=x,
                y=y,
                width=width,
                height=height,
                options=list(options),
                value=str(options[0]),
            )
        else:
            tooltip = "Data (DD/MM/AAAA)" if field_type == "date" else None
            c.acroForm.textfield(
                name=field_name,
                x=x,
                y=y,
                width=width,
                height=height,
                borderStyle="inset",
                tooltip=tooltip,
            )
        c.showPage()
        c.save()
        buffer.seek(0)

        overlay_reader = PdfReader(buffer)
        overlay_page = overlay_reader.pages[0]
        overlay_acroform = cast(DictionaryObject, overlay_reader.root_object["/AcroForm"].get_object())

        annots_before = list(target_page.get("/Annots", []))
        target_page.merge_page(overlay_page)
        annots_after = list(target_page.get("/Annots", []))
        # Comparação por identidade (não por posição): merge_page normalmente só anexa
        # ao final de /Annots, mas depender dessa ordem seria uma suposição implícita
        # sobre o comportamento interno do pypdf.
        new_annots = [annot for annot in annots_after if annot not in annots_before]

        acroform = writer.root_object.get("/AcroForm")
        if acroform is None:
            acroform = writer._add_object(overlay_acroform.clone(writer))
            acroform_obj = cast(DictionaryObject, acroform.get_object())
            acroform_obj[NameObject("/Fields")] = ArrayObject(new_annots)
            writer.root_object[NameObject("/AcroForm")] = acroform
        else:
            acroform_obj = cast(DictionaryObject, acroform.get_object())
            fields = acroform_obj.get("/Fields", ArrayObject())
            for annot in new_annots:
                fields.append(annot)
            acroform_obj[NameObject("/Fields")] = fields

            _merge_default_fonts(writer, acroform_obj, overlay_acroform)

        acroform_obj[NameObject("/NeedAppearances")] = BooleanObject(True)
        write_pdf(writer, output_path)


def _require_new_field_name(reader: PdfReader, field_name: str) -> None:
    """Dois campos raiz com o mesmo nome deixam o formulário ambíguo: cada leitor decide de
    um jeito qual valor vale. Por isso o nome precisa ser novo."""
    if field_name in (reader.get_fields() or {}):
        raise ValueError(Message('Já existe um campo chamado "{name}" neste PDF. Escolha outro nome.', name=field_name))


def _merge_default_fonts(writer: PdfWriter, acroform_obj: DictionaryObject, overlay_acroform) -> None:
    """Registra no `/DR` do formulário as fontes que o campo novo usa (`/DA`). Sem isso, num
    formulário que ainda não tinha `/DR`, o leitor não acha a fonte para desenhar o valor."""
    overlay_fonts = overlay_acroform.get_object().get("/DR", {}).get("/Font", {})
    if not overlay_fonts:
        return
    resources = acroform_obj.get("/DR")
    if resources is None:
        resources = DictionaryObject()
        acroform_obj[NameObject("/DR")] = resources
    resources = cast(DictionaryObject, resources.get_object())
    fonts = resources.get("/Font")
    if fonts is None:
        fonts = DictionaryObject()
        resources[NameObject("/Font")] = fonts
    fonts = cast(DictionaryObject, fonts.get_object())
    for font_name, font_ref in overlay_fonts.items():
        if font_name not in fonts:
            fonts[NameObject(font_name)] = font_ref.clone(writer)


@dataclass
class RadioOption:
    """Uma opção de um grupo de botões de opção (radio): valor exportado e posição."""

    value: str
    page_number: int
    x: float
    y: float


class AddRadioGroup(PDFOperation):
    """Adiciona um grupo de botões de opção (radio) — várias opções mutuamente exclusivas
    compartilhando um único campo, cada uma podendo estar numa página diferente (comum em
    formulários longos com "assinale uma opção" repetido em várias páginas).

    Diferente de um campo de texto/caixa de seleção (onde o widget na página *é* o campo),
    um grupo radio no PDF é um campo-pai (`/FT /Btn`) sem posição própria, com um widget
    filho (`/Kids`) por opção. O reportlab já monta essa estrutura corretamente ao desenhar
    todas as opções num único `Canvas` com o mesmo `name` — o que exigiu confirmação
    empírica nesta sessão foi que juntar os widgets de páginas diferentes de volta ao
    documento original via `merge_page` (a mesma técnica de `AddFormField`) preserva o
    vínculo `/Parent`/`/Kids` só se ele for reconstruído à mão depois do merge: o pypdf
    clona cada widget isoladamente por chamada de `merge_page`, então o campo-pai
    compartilhado precisa ser criado uma única vez aqui e referenciado por todos os widgets
    já clonados, em vez de confiar em `merge_page` para propagar essa relação sozinho.
    """

    def run(
        self,
        input_path: str,
        output_path: str,
        field_name: str,
        options: list[RadioOption],
        size: float = 16,
        selected_value: str | None = None,
    ) -> None:
        if len(options) < 2:
            raise ValueError("Um grupo de opções precisa de pelo menos duas opções.")
        values = [opt.value for opt in options]
        if len(set(values)) != len(values):
            raise ValueError("Os valores das opções devem ser únicos.")
        if selected_value is not None and selected_value not in values:
            raise ValueError("selected_value deve ser um dos valores das opções.")

        reader = open_reader(input_path)
        for opt in options:
            require_valid_page_index(len(reader.pages), opt.page_number, label="Página em options")
        _require_new_field_name(reader, field_name)

        writer = PdfWriter(clone_from=reader)

        options_by_page: dict[int, list[RadioOption]] = {}
        for opt in options:
            options_by_page.setdefault(opt.page_number, []).append(opt)
        max_page = max(options_by_page)

        buffer = io.BytesIO()
        c = canvas.Canvas(buffer)
        for page_index in range(max_page + 1):
            page_width = float(writer.pages[page_index].mediabox.width)
            page_height = float(writer.pages[page_index].mediabox.height)
            c.setPageSize((page_width, page_height))
            for opt in options_by_page.get(page_index, []):
                c.acroForm.radio(
                    name=field_name,
                    value=opt.value,
                    selected=(opt.value == selected_value),
                    x=opt.x,
                    y=opt.y,
                    size=size,
                )
            c.showPage()
        c.save()
        buffer.seek(0)

        overlay_reader = PdfReader(buffer)
        overlay_acroform = cast(DictionaryObject, overlay_reader.root_object["/AcroForm"].get_object())
        overlay_parent = cast(DictionaryObject, cast(ArrayObject, overlay_acroform["/Fields"])[0].get_object())

        kid_refs: list[IndirectObject] = []
        for page_index in options_by_page:
            target_page = writer.pages[page_index]
            annots_before = list(target_page.get("/Annots", []))
            target_page.merge_page(overlay_reader.pages[page_index])
            annots_after = list(target_page.get("/Annots", []))
            kid_refs.extend(annot for annot in annots_after if annot not in annots_before)

        parent = DictionaryObject()
        parent[NameObject("/FT")] = NameObject("/Btn")
        parent[NameObject("/Ff")] = overlay_parent.raw_get("/Ff")
        parent[NameObject("/T")] = overlay_parent.raw_get("/T")
        if "/V" in overlay_parent:
            parent[NameObject("/V")] = overlay_parent.raw_get("/V")
        parent[NameObject("/Kids")] = ArrayObject(kid_refs)
        parent_ref = writer._add_object(parent)
        for kid_ref in kid_refs:
            cast(DictionaryObject, kid_ref.get_object())[NameObject("/Parent")] = parent_ref

        acroform = writer.root_object.get("/AcroForm")
        if acroform is None:
            acroform = writer._add_object(overlay_acroform.clone(writer))
            acroform_obj = cast(DictionaryObject, acroform.get_object())
            acroform_obj[NameObject("/Fields")] = ArrayObject([parent_ref])
            writer.root_object[NameObject("/AcroForm")] = acroform
        else:
            acroform_obj = cast(DictionaryObject, acroform.get_object())
            fields = acroform_obj.get("/Fields", ArrayObject())
            fields.append(parent_ref)
            acroform_obj[NameObject("/Fields")] = fields

            _merge_default_fonts(writer, acroform_obj, overlay_acroform)

        acroform_obj[NameObject("/NeedAppearances")] = BooleanObject(True)
        write_pdf(writer, output_path)


@dataclass
class FormFieldInfo:
    """Descrição resumida de um campo de formulário já existente num PDF."""

    name: str
    field_type: str
    value: str


# Bits de /Ff (ISO 32000-1, tabela 221/226/227) que distinguem subtipos de /Btn e /Ch.
_FLAG_RADIO = 1 << 15
_FLAG_PUSHBUTTON = 1 << 16
_FLAG_COMBO = 1 << 17


def _describe_field_type(field_type, flags) -> str:
    flags = int(flags or 0)
    if field_type == "/Tx":
        return "text"
    if field_type == "/Btn":
        if flags & _FLAG_RADIO:
            return "radio"
        if flags & _FLAG_PUSHBUTTON:
            return "button"
        return "checkbox"
    if field_type == "/Ch":
        return "dropdown" if flags & _FLAG_COMBO else "listbox"
    if field_type == "/Sig":
        return "signature"
    return "unknown"


class ListFormFields(PDFOperation):
    """Lista os campos de formulário já existentes num PDF: nome, tipo e valor atual.
    Campos que só agrupam outros (nomes hierárquicos como "cliente" em "cliente.nome") não
    são campos preenchíveis e ficam de fora."""

    def run(self, input_path: str) -> list[FormFieldInfo]:
        reader = open_reader(input_path, purpose=READ)
        fields = reader.get_fields() or {}
        return [
            FormFieldInfo(
                name=name,
                field_type=_describe_field_type(field.field_type, field.flags),
                value="" if field.value is None else str(field.value),
            )
            for name, field in fields.items()
            if field.field_type is not None
        ]


_FLAG_READ_ONLY = 1


@dataclass
class FormWidget:
    """Um widget de campo numa página, como o visualizador precisa para preenchê-lo."""

    page_index: int
    name: str  # nome completo do campo
    field_type: str  # como em FormFieldInfo
    rect: tuple[float, float, float, float]  # espaço do usuário do PDF
    value: str
    on_state: str = ""  # caixa de seleção/opção: o estado "ligado" deste widget ("/Yes", "/M")
    options: list[tuple[str, str]] = field(default_factory=list)  # lista: (valor, texto)
    read_only: bool = False


def _inherited(node, key: str):
    """Valor de `key` no widget ou no primeiro campo-pai que o tenha (/FT, /Ff, /V, /Opt)."""
    visited: set[int] = set()
    while node is not None and id(node) not in visited:
        visited.add(id(node))
        if key in node:
            return node[key]
        parent = node.get("/Parent")
        node = parent.get_object() if parent is not None else None
    return None


def _page_widgets(page):
    for annotation in page["/Annots"] if "/Annots" in page else []:
        annotation = annotation.get_object()
        if annotation.get("/Subtype") == "/Widget":
            yield annotation


def _appearance_states(widget) -> list[str]:
    appearance = widget.get("/AP")
    normal = appearance.get_object().get("/N") if appearance is not None else None
    normal = normal.get_object() if normal is not None else None
    # Num campo de texto /N é a própria aparência (um stream), não um dicionário de estados.
    if not isinstance(normal, DictionaryObject) or isinstance(normal, StreamObject):
        return []
    return [str(key) for key in normal]


def _choice_options(widget) -> list[tuple[str, str]]:
    options = []
    for option in _inherited(widget, "/Opt") or []:
        option = option.get_object()
        if isinstance(option, ArrayObject) and len(option) >= 2:
            options.append((str(option[0]), str(option[1])))
        else:
            options.append((str(option), str(option)))
    return options


class ListFormWidgets(PDFOperation):
    """Os widgets preenchíveis de cada página (texto, caixa de seleção, opção e listas),
    com posição e valor atual. Botões e assinaturas ficam de fora."""

    def run(self, input_path: str) -> list[FormWidget]:
        reader = open_reader(input_path, purpose=READ)
        if "/AcroForm" not in reader.root_object:
            return []
        widgets = []
        for page_index, page in enumerate(reader.pages):
            for widget in _page_widgets(page):
                flags = int(_inherited(widget, "/Ff") or 0)
                field_type = _describe_field_type(_inherited(widget, "/FT"), flags)
                name = _qualified_name(widget)
                if field_type not in ("text", "checkbox", "radio", "dropdown", "listbox") or not name:
                    continue
                value = _inherited(widget, "/V")
                left, bottom, right, top = (float(v) for v in widget["/Rect"])
                states = [state for state in _appearance_states(widget) if state != "/Off"]
                widgets.append(
                    FormWidget(
                        page_index=page_index,
                        name=name,
                        field_type=field_type,
                        rect=(min(left, right), min(bottom, top), max(left, right), max(bottom, top)),
                        value="" if value is None else str(value),
                        on_state=states[0] if states else "",
                        options=_choice_options(widget),
                        read_only=bool(flags & _FLAG_READ_ONLY),
                    )
                )
        return widgets


class FillFormFields(PDFOperation):
    """Grava valores nos campos, pelo nome completo. Caixas de seleção e opções recebem o
    estado ("/Yes", "/Off", "/M"); os outros campos, o texto ou o valor da opção, com a
    aparência refeita pelo pypdf."""

    def run(self, input_path: str, output_path: str, values: dict[str, str]) -> None:
        writer = PdfWriter(clone_from=open_reader(input_path))
        texts = {}
        for page in writer.pages:
            for widget in _page_widgets(page):
                name = _qualified_name(widget)
                if name not in values:
                    continue
                if _inherited(widget, "/FT") != "/Btn":
                    texts[name] = values[name]
                    continue
                # O pypdf grava o /V de um grupo de opções como texto, não como nome; aqui
                # cada widget liga só se tiver a aparência desse estado.
                state = values[name]
                widget[NameObject("/AS")] = NameObject(state if state in _appearance_states(widget) else "/Off")
                field_node = widget
                while "/T" not in field_node:
                    field_node = field_node["/Parent"].get_object()
                field_node[NameObject("/V")] = NameObject(state)
        if texts:
            writer.update_page_form_field_values(None, texts, auto_regenerate=False)
        write_pdf(writer, output_path)


class ExportFormDataCSV(PDFOperation):
    """Exporta nome, tipo e valor atual de cada campo de um PDF para CSV (ou Excel, se o
    arquivo terminar em .xlsx). Ver `core/tabular.py`: UTF-8 com BOM e valores que o Excel
    leria como fórmula protegidos."""

    def run(self, input_path: str, output_path: str) -> int:
        fields = ListFormFields().run(input_path)
        write_table(output_path, ["campo", "tipo", "valor"], ([f.name, f.field_type, f.value] for f in fields))
        return len(fields)


@dataclass
class FormImportResult:
    """Resultado de uma importação de dados: quantos campos foram preenchidos e quais nomes
    do arquivo de dados não existem no PDF."""

    filled: int
    not_found: list[str] = field(default_factory=list)


def _read_csv_text(csv_path: str) -> str:
    raw = Path(csv_path).read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _read_csv_values(csv_path: str) -> dict[str, str]:
    """Lê "campo"/"valor" (ou "name"/"value") de um CSV, aceitando o que o Excel grava em
    português: UTF-8 com BOM ou Windows-1252, separado por vírgula ou ponto e vírgula. Um
    arquivo .xlsx (o que `ExportFormDataCSV` grava com esse nome) também serve."""
    if Path(csv_path).suffix.lower() == ".xlsx":
        rows = _read_xlsx_rows(csv_path)
    else:
        text = _read_csv_text(csv_path)
        # O Sniffer só escolhe o separador: as aspas são sempre as duplas do padrão. Deixado
        # livre, ele confunde o apóstrofo que protege fórmulas com o caractere de aspas.
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    values: dict[str, str] = {}
    for row in rows:
        normalized = {str(key).strip().lower(): (value or "") for key, value in row.items() if key is not None}
        name = _unescape_formula(str(normalized.get("campo") or normalized.get("name") or "").strip())
        if not name:
            continue
        values[name] = _unescape_formula(str(normalized.get("valor") or normalized.get("value") or ""))
    return values


def _unescape_formula(text: str) -> str:
    """Desfaz o apóstrofo que a exportação põe na frente de valores que o Excel leria como
    fórmula (ver `core/tabular.py`)."""
    if text.startswith("'") and escape_formula(text[1:]) == text:
        return text[1:]
    return text


def _read_xlsx_rows(path: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    header = ["" if value is None else str(value) for value in rows[0]]
    return [dict(zip(header, ("" if value is None else str(value) for value in row))) for row in rows[1:]]


def _fill_fields(reader: PdfReader, output_path: str, values: dict[str, str]) -> FormImportResult:
    """Preenche os campos que existem no PDF e informa os que não existem. Se nenhum nome
    do arquivo de dados existir, é erro: relatar "preenchido" sem ter preenchido nada
    esconderia um arquivo de dados trocado ou com cabeçalho errado."""
    fields = reader.get_fields() or {}
    known = set(fields) | {name.rsplit(".", 1)[-1] for name in fields}
    matched = {name: value for name, value in values.items() if name in known}
    not_found = sorted(name for name in values if name not in known)
    if not matched:
        raise ValueError(
            Message(
                "Nenhum dos campos do arquivo de dados existe neste PDF: {names}",
                names=", ".join(not_found[:10]),
            )
        )
    writer = PdfWriter(clone_from=reader)
    writer.update_page_form_field_values(list(writer.pages), matched, auto_regenerate=False)
    write_pdf(writer, output_path)
    return FormImportResult(filled=len(matched), not_found=not_found)


class ImportFormDataCSV(PDFOperation):
    """Preenche os campos de um PDF a partir de um CSV com colunas "campo" e "valor"
    (aceita também "name"/"value", para CSVs exportados de outras ferramentas). Caixas de
    seleção e listas recebem o valor bruto do PDF (ex.: "/Yes", "/Off", o nome de uma
    opção) — não um rótulo amigável — porque é esse o mesmo formato que `ExportFormDataCSV`
    grava, então uma volta completa exportar → editar → importar sempre é consistente."""

    def run(self, input_path: str, output_path: str, csv_path: str) -> FormImportResult:
        reader = open_reader(input_path)
        values = _read_csv_values(csv_path)
        if not values:
            raise ValueError("Nenhum campo com nome válido encontrado no CSV.")
        return _fill_fields(reader, output_path, values)


def _fdf_hex_string(value: str) -> str:
    # Hexadecimal com BOM UTF-16BE (em vez de string literal com parênteses) evita ter que
    # escapar "\\", "(" e ")" na mão — a fonte mais comum de FDF malformado feito à unha — e
    # já suporta acentos/Unicode sem depender de PDFDocEncoding.
    return "<FEFF" + value.encode("utf-16-be").hex().upper() + ">"


def _fdf_name(value: str) -> str:
    """Um nome PDF (/Yes, /Opção_1) com os caracteres especiais escapados como #xx."""
    buffer = io.BytesIO()
    NameObject("/" + value.lstrip("/")).write_to_stream(buffer)
    return buffer.getvalue().decode("latin-1")


def _fdf_field_lines(name: str, node: dict, lines: list[str]) -> None:
    lines.append("<<")
    lines.append(f"/T {_fdf_hex_string(name)}")
    if "value" in node:
        info = node["value"]
        if info.field_type in ("checkbox", "radio") and info.value.startswith("/"):
            # Caixas de seleção e grupos de opções guardam o estado como nome (/Yes, /Off);
            # gravar como texto ("/Yes") faz o Acrobat ignorar o valor ao importar.
            lines.append(f"/V {_fdf_name(info.value)}")
        else:
            lines.append(f"/V {_fdf_hex_string(info.value)}")
    if node["kids"]:
        lines.append("/Kids [")
        for kid_name, kid in node["kids"].items():
            _fdf_field_lines(kid_name, kid, lines)
        lines.append("]")
    lines.append(">>")


class ExportFormDataFDF(PDFOperation):
    """Exporta os valores atuais dos campos de um PDF para um arquivo FDF (o formato nativo
    do Acrobat para dados de formulário) — útil para reaproveitar os mesmos dados em outra
    cópia do mesmo formulário. Nomes hierárquicos ("cliente.nome") viram campos aninhados
    em `/Kids`, como o Acrobat grava."""

    def run(self, input_path: str, output_path: str) -> int:
        fields = ListFormFields().run(input_path)
        tree: dict = {"kids": {}}
        for info in fields:
            node = tree
            for part in info.name.split("."):
                node = node["kids"].setdefault(part, {"kids": {}})
            node["value"] = info
        lines = ["%FDF-1.2", "1 0 obj", "<<", "/FDF", "<<", "/Fields ["]
        for name, node in tree["kids"].items():
            _fdf_field_lines(name, node, lines)
        lines.extend(["]", ">>", ">>", "endobj", "trailer", "<<", "/Root 1 0 R", ">>", "%%EOF"])
        with atomic_output(output_path) as temp_path:
            with open(temp_path, "wb") as f:
                f.write("\n".join(lines).encode("latin-1"))
        return len(fields)


def _fdf_value(value) -> str | None:
    value = value.get_object() if hasattr(value, "get_object") and not isinstance(value, (str, bytes)) else value
    if isinstance(value, NameObject):
        return str(value)  # "/Yes": o mesmo formato bruto que o CSV usa
    if isinstance(value, ArrayObject):
        return _fdf_value(value[0]) if value else ""  # lista com várias seleções: a primeira
    if isinstance(value, (str, bytes)):
        return value if isinstance(value, str) else value.decode("latin-1")
    return None


def _parse_fdf(data: bytes) -> dict[str, str]:
    """Lê os campos de um FDF com o parser de objetos do próprio pypdf: strings literais
    com parênteses aninhados e escapes (\\n, \\ddd), hexadecimais, nomes (/Yes) e campos
    aninhados em /Kids (o nome completo junta as partes com ".")."""
    marker = data.find(b"/Fields")
    if marker < 0:
        return {}
    start = marker + len(b"/Fields")
    while start < len(data) and data[start:start + 1].isspace():
        start += 1
    stream = io.BytesIO(data)
    stream.seek(start)
    try:
        fields = read_object(stream, None)
    except Exception as exc:
        raise ValueError(Message("O arquivo FDF está malformado: {error}", error=exc)) from exc
    values: dict[str, str] = {}

    def visit(entries, prefix: str, depth: int = 0) -> None:
        if not isinstance(entries, ArrayObject) or depth > 32:
            return
        for entry in entries:
            if not isinstance(entry, DictionaryObject):
                continue  # referência indireta: o FDF não é um PDF completo para resolvê-la
            part = _fdf_value(entry.get("/T"))
            name = f"{prefix}.{part}" if prefix and part else (part or prefix)
            if "/V" in entry and name:
                value = _fdf_value(entry["/V"])
                if value is not None:
                    values[name] = value
            visit(entry.get("/Kids"), name or "", depth + 1)

    visit(fields, "")
    return values


class ImportFormDataFDF(PDFOperation):
    """Importa valores de campos a partir de um arquivo FDF e preenche o PDF.

    Entende os FDFs gravados pelo Acrobat e por `ExportFormDataFDF`: valores de texto
    (literais ou hexadecimais, com escapes), estados de caixas de seleção como nome (/Yes) e
    nomes hierárquicos em `/Kids`. Numa lista com várias seleções, vale a primeira.
    """

    def run(self, input_path: str, output_path: str, fdf_path: str) -> FormImportResult:
        reader = open_reader(input_path)
        with open(fdf_path, "rb") as f:
            values = _parse_fdf(f.read())
        if not values:
            raise ValueError("Nenhum campo encontrado no arquivo FDF.")
        return _fill_fields(reader, output_path, values)


def _qualified_name(widget) -> str:
    """Nome completo do campo de um widget ("cliente.nome"), subindo pelos `/Parent`."""
    parts: list[str] = []
    node = widget.get_object()
    visited: set[int] = set()
    while node is not None and id(node) not in visited:
        visited.add(id(node))
        if "/T" in node:
            parts.insert(0, str(node["/T"]))
        parent = node.get("/Parent")
        node = parent.get_object() if parent is not None else None
    return ".".join(parts)


class FlattenFormFields(PDFOperation):
    """Achata os campos de formulário de um PDF: grava o valor atual de cada campo
    preenchido como conteúdo estático da página e remove o widget interativo — um "Fill &
    Sign" simplificado, para garantir que os valores não sejam alterados depois de
    preenchidos. Campos sem valor definido são deixados como estão (não há o que achatar
    num campo em branco).

    Vale para qualquer formulário, inclusive os com nomes hierárquicos ("cliente.nome",
    comuns em formulários do Acrobat): o widget é identificado pelo nome completo do campo,
    e os campos que ficam sem widget saem da árvore do formulário.
    """

    def run(self, input_path: str, output_path: str) -> int:
        reader = open_reader(input_path)
        writer = PdfWriter(clone_from=reader)

        fields = writer.get_fields() or {}
        values = {
            name: info.value
            for name, info in fields.items()
            if info.value is not None and info.field_type is not None
        }
        if values:
            writer.update_page_form_field_values(
                list(writer.pages), values, auto_regenerate=False, flatten=True
            )
            flattened_widgets = []
            for page in writer.pages:
                if "/Annots" not in page:
                    continue
                kept = ArrayObject()
                for ref in cast(ArrayObject, page["/Annots"]):
                    annotation = ref.get_object()
                    if annotation.get("/Subtype") == "/Widget" and _qualified_name(annotation) in values:
                        flattened_widgets.append(ref)
                    else:
                        kept.append(ref)
                if kept:
                    page[NameObject("/Annots")] = kept
                else:
                    del page["/Annots"]
            remove_widgets_from_form(writer, flattened_widgets)
            acroform = writer.root_object.get("/AcroForm")
            if acroform is not None and not acroform.get_object().get("/Fields"):
                del writer.root_object["/AcroForm"]

        write_pdf(writer, output_path)
        return len(values)
