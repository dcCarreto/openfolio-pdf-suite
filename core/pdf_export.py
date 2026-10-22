"""Conversão de PDF para Word e Excel (A09), a partir do texto das páginas.

Só o texto vai: imagens, cores e a posição exata ficam de fora, e página escaneada sem OCR
não tem texto nenhum (as operações devolvem essas páginas para a tela avisar).

- Word: o texto corrido de cada página (`extract_text`), com as linhas reunidas em
  parágrafos — uma linha bem mais curta que as outras fecha o parágrafo (é a última linha
  dele, ou um título) — e uma quebra de página entre as páginas.
- Excel: uma planilha por página, a partir do texto com layout (`extraction_mode="layout"`),
  que mantém as colunas alinhadas. Trechos separados por dois espaços ou mais são células;
  cada célula vai para a coluna cuja posição ela ocupa na página, então uma célula vazia
  no meio da linha não desloca as seguintes. Números ("980", "1.234,56", "12.5") viram
  números; o resto vira texto, nunca fórmula.
"""

import bisect
import re

from .base import EXTRACT, atomic_output, open_reader
from .extract_text import page_text
from .progress import Progress, ensure_progress

_CELL = re.compile(r"\S+(?: \S+)*")
_SHORT_LINE = 0.75  # fração da linha mais longa da página abaixo da qual a linha fecha o parágrafo
# Caracteres de controle que alguns PDFs trazem no texto: o XML do .docx e do .xlsx os recusa.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean_text(page, layout: bool = False) -> str:
    return _CONTROL.sub("", page_text(page, layout)).replace("\t", " ")


def paragraphs(text: str) -> list[str]:
    """As linhas de uma página reunidas em parágrafos."""
    lines = [line.strip() for line in text.splitlines()]
    longest = max((len(line) for line in lines), default=0)
    result, current = [], ""
    for line in lines:
        if not line:
            if current:
                result.append(current)
            current = ""
            continue
        if current.endswith("-"):  # palavra hifenizada no fim da linha
            current = current[:-1] + line
        else:
            current = f"{current} {line}" if current else line
        if len(line) < _SHORT_LINE * longest:
            result.append(current)
            current = ""
    if current:
        result.append(current)
    return result


def table_rows(layout_text: str) -> list[list[str]]:
    """As linhas do texto com layout de uma página, cada uma dividida nas colunas da página."""
    lines = [[(match.start(), match.group()) for match in _CELL.finditer(line)] for line in layout_text.splitlines()]
    lines = [cells for cells in lines if cells]
    anchors: list[int] = []
    for start in sorted({start for cells in lines if len(cells) > 1 for start, _text in cells}):
        if not anchors or start - anchors[-1] > 2:
            anchors.append(start)
    rows = []
    for cells in lines:
        if len(cells) == 1:  # texto corrido: fica na primeira coluna
            rows.append([cells[0][1]])
            continue
        row: list[str] = []
        for start, text in cells:
            column = max(bisect.bisect_right(anchors, start) - 1, len(row))
            row.extend([""] * (column - len(row)))
            row.append(text)
        rows.append(row)
    return rows


def cell_value(text: str):
    """Número quando o texto é um número sem ambiguidade; senão, o próprio texto.

    "1.234" fica texto (mil e duzentos ou um vírgula dois?), assim como números com zero à
    esquerda (códigos, CEP) ou longos demais para o Excel guardar sem perder dígitos.
    """
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})", text):
        return int(text)
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*,\d+|-?\d+,\d+", text):
        return float(text.replace(".", "").replace(",", "."))
    if re.fullmatch(r"-?\d+\.\d+", text) and not re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):
        return float(text)
    return text


class ConvertPDFToWord:
    def run(self, input_path: str, output_path: str, progress: Progress | None = None) -> list[int]:
        """Grava o .docx e devolve as páginas (a partir de 0) sem texto."""
        from docx import Document

        progress = ensure_progress(progress)
        reader = open_reader(input_path, purpose=EXTRACT)
        document = Document()
        empty = []
        for index, page in enumerate(progress.steps(reader.pages, message="Convertendo páginas")):
            if index:
                document.add_page_break()
            found = paragraphs(_clean_text(page))
            if not found:
                empty.append(index)
            for paragraph in found:
                document.add_paragraph(paragraph)
        with atomic_output(output_path) as temp_path:
            document.save(temp_path)
        return empty


class ConvertPDFToExcel:
    def run(self, input_path: str, output_path: str, progress: Progress | None = None) -> list[int]:
        """Grava o .xlsx (uma planilha por página) e devolve as páginas sem texto."""
        from openpyxl import Workbook

        progress = ensure_progress(progress)
        reader = open_reader(input_path, purpose=EXTRACT)
        workbook = Workbook()
        workbook.remove(workbook.active)
        empty = []
        for index, page in enumerate(progress.steps(reader.pages, message="Convertendo páginas")):
            sheet = workbook.create_sheet(str(index + 1))  # o número da página, em qualquer idioma
            rows = table_rows(_clean_text(page, layout=True))
            if not rows:
                empty.append(index)
            for row in rows:
                sheet.append([cell_value(text) for text in row])
            for cells in sheet.iter_rows():
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"  # texto, nunca fórmula (o PDF vem de terceiros)
        with atomic_output(output_path) as temp_path:
            workbook.save(temp_path)
        return empty
