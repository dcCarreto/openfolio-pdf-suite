import pytest
from pypdf import PdfReader, PdfWriter

from core.base import EncryptedPDFError, add_object
from core.create import CreateBlankPDF
from core.form_fields import (
    AddFormField,
    AddRadioGroup,
    ExportFormDataCSV,
    ExportFormDataFDF,
    FlattenFormFields,
    ImportFormDataCSV,
    ImportFormDataFDF,
    ListFormFields,
    RadioOption,
)


def test_add_text_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_field.pdf"
    AddFormField().run(
        str(pdf_path),
        str(output_path),
        field_name="nome",
        page_number=0,
        x=50,
        y=300,
        width=150,
        height=20,
        field_type="text",
    )

    fields = PdfReader(str(output_path)).get_fields()
    assert "nome" in fields
    assert fields["nome"]["/FT"] == "/Tx"


def test_add_checkbox_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_checkbox.pdf"
    AddFormField().run(
        str(pdf_path),
        str(output_path),
        field_name="aceite",
        page_number=0,
        x=50,
        y=250,
        height=15,
        field_type="checkbox",
        checked=True,
    )

    fields = PdfReader(str(output_path)).get_fields()
    assert "aceite" in fields
    assert fields["aceite"]["/FT"] == "/Btn"


def test_add_two_fields_sequentially_keeps_both(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    step1_path = tmp_path / "step1.pdf"
    AddFormField().run(
        str(pdf_path),
        str(step1_path),
        field_name="campo1",
        page_number=0,
        x=50,
        y=300,
        field_type="text",
    )

    step2_path = tmp_path / "step2.pdf"
    AddFormField().run(
        str(step1_path),
        str(step2_path),
        field_name="campo2",
        page_number=0,
        x=50,
        y=250,
        height=15,
        field_type="checkbox",
    )

    fields = PdfReader(str(step2_path)).get_fields()
    assert set(fields.keys()) == {"campo1", "campo2"}


def test_rejects_invalid_field_type(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddFormField().run(
            str(pdf_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            page_number=0,
            x=0,
            y=0,
            field_type="invalido",
        )


def test_rejects_out_of_range_page_number(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddFormField().run(
            str(pdf_path), str(tmp_path / "out.pdf"), field_name="x", page_number=5, x=0, y=0
        )


def test_rejects_encrypted_input(make_encrypted_pdf, tmp_path):
    encrypted_path = make_encrypted_pdf("protected.pdf", [(200, 200)])

    with pytest.raises(EncryptedPDFError):
        AddFormField().run(
            str(encrypted_path), str(tmp_path / "out.pdf"), field_name="x", page_number=0, x=0, y=0
        )


def test_add_date_field_is_a_text_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_date.pdf"
    AddFormField().run(
        str(pdf_path), str(output_path), field_name="nascimento", page_number=0, x=50, y=300, field_type="date"
    )

    fields = PdfReader(str(output_path)).get_fields()
    assert fields["nascimento"]["/FT"] == "/Tx"


def test_add_dropdown_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_dropdown.pdf"
    AddFormField().run(
        str(pdf_path),
        str(output_path),
        field_name="estado",
        page_number=0,
        x=50,
        y=300,
        field_type="dropdown",
        options=["SP", "RJ", "MG"],
    )

    fields = ListFormFields().run(str(output_path))
    field = next(f for f in fields if f.name == "estado")
    assert field.field_type == "dropdown"


def test_add_listbox_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_listbox.pdf"
    AddFormField().run(
        str(pdf_path),
        str(output_path),
        field_name="opcoes",
        page_number=0,
        x=50,
        y=300,
        field_type="listbox",
        options=["A", "B", "C"],
    )

    fields = ListFormFields().run(str(output_path))
    field = next(f for f in fields if f.name == "opcoes")
    assert field.field_type == "listbox"


def test_choice_field_without_options_raises(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddFormField().run(
            str(pdf_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            page_number=0,
            x=0,
            y=0,
            field_type="dropdown",
        )


def test_add_radio_group_on_single_page(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    output_path = tmp_path / "with_radio.pdf"
    AddRadioGroup().run(
        str(pdf_path),
        str(output_path),
        field_name="sexo",
        options=[
            RadioOption(value="M", page_number=0, x=50, y=300),
            RadioOption(value="F", page_number=0, x=100, y=300),
        ],
        selected_value="F",
    )

    reader = PdfReader(str(output_path))
    fields = reader.get_fields()
    assert fields["sexo"]["/FT"] == "/Btn"
    assert fields["sexo"]["/V"] == "/F"

    root_field = reader.trailer["/Root"]["/AcroForm"]["/Fields"][0].get_object()
    assert len(root_field["/Kids"]) == 2
    for kid_ref in root_field["/Kids"]:
        assert kid_ref.get_object()["/Parent"].get_object() is root_field


def test_add_radio_group_across_multiple_pages(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=2, width=400, height=400)

    output_path = tmp_path / "with_radio.pdf"
    AddRadioGroup().run(
        str(pdf_path),
        str(output_path),
        field_name="opcao",
        options=[
            RadioOption(value="sim", page_number=0, x=50, y=300),
            RadioOption(value="nao", page_number=1, x=50, y=300),
        ],
    )

    reader = PdfReader(str(output_path))
    root_field = reader.trailer["/Root"]["/AcroForm"]["/Fields"][0].get_object()
    assert len(root_field["/Kids"]) == 2
    page0_annots = list(reader.pages[0].get("/Annots", []))
    page1_annots = list(reader.pages[1].get("/Annots", []))
    assert len(page0_annots) == 1
    assert len(page1_annots) == 1


def test_add_radio_group_requires_at_least_two_options(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddRadioGroup().run(
            str(pdf_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            options=[RadioOption(value="A", page_number=0, x=0, y=0)],
        )


def test_add_radio_group_rejects_duplicate_values(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddRadioGroup().run(
            str(pdf_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            options=[
                RadioOption(value="A", page_number=0, x=0, y=0),
                RadioOption(value="A", page_number=0, x=50, y=0),
            ],
        )


def test_add_radio_group_rejects_invalid_selected_value(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)

    with pytest.raises(ValueError):
        AddRadioGroup().run(
            str(pdf_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            options=[
                RadioOption(value="A", page_number=0, x=0, y=0),
                RadioOption(value="B", page_number=0, x=50, y=0),
            ],
            selected_value="C",
        )


def test_add_radio_group_rejects_encrypted_input(make_encrypted_pdf, tmp_path):
    encrypted_path = make_encrypted_pdf("protected.pdf", [(200, 200)])

    with pytest.raises(EncryptedPDFError):
        AddRadioGroup().run(
            str(encrypted_path),
            str(tmp_path / "out.pdf"),
            field_name="x",
            options=[
                RadioOption(value="A", page_number=0, x=0, y=0),
                RadioOption(value="B", page_number=0, x=50, y=0),
            ],
        )


def _build_pdf_with_various_fields(tmp_path) -> str:
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)

    step1 = tmp_path / "step1.pdf"
    AddFormField().run(
        str(pdf_path), str(step1), field_name="nome", page_number=0, x=50, y=350, field_type="text"
    )
    step2 = tmp_path / "step2.pdf"
    AddFormField().run(
        str(step1), str(step2), field_name="aceite", page_number=0, x=50, y=300, height=15, field_type="checkbox"
    )
    step3 = tmp_path / "step3.pdf"
    AddFormField().run(
        str(step2),
        str(step3),
        field_name="estado",
        page_number=0,
        x=50,
        y=250,
        field_type="dropdown",
        options=["SP", "RJ"],
    )
    final_path = tmp_path / "final.pdf"
    AddRadioGroup().run(
        str(step3),
        str(final_path),
        field_name="sexo",
        options=[
            RadioOption(value="M", page_number=0, x=50, y=200),
            RadioOption(value="F", page_number=0, x=100, y=200),
        ],
        selected_value="M",
    )
    return str(final_path)


def test_list_form_fields_describes_each_type(tmp_path):
    pdf_path = _build_pdf_with_various_fields(tmp_path)

    fields = {f.name: f.field_type for f in ListFormFields().run(pdf_path)}

    assert fields == {
        "nome": "text",
        "aceite": "checkbox",
        "estado": "dropdown",
        "sexo": "radio",
    }


def test_export_and_import_csv_round_trip(tmp_path):
    pdf_path = _build_pdf_with_various_fields(tmp_path)
    csv_path = tmp_path / "dados.csv"

    count = ExportFormDataCSV().run(pdf_path, str(csv_path))
    assert count == 4

    content = csv_path.read_text(encoding="utf-8")
    assert "nome" in content

    # edita o CSV exportado para simular um preenchimento em lote
    lines = content.splitlines()
    lines = [line.replace("nome,text,", "nome,text,João da Silva") if line.startswith("nome,") else line for line in lines]
    csv_path.write_text("\n".join(lines), encoding="utf-8")

    output_path = tmp_path / "preenchido.pdf"
    imported = ImportFormDataCSV().run(pdf_path, str(output_path), str(csv_path))
    assert imported.filled == 4 and imported.not_found == []

    fields = PdfReader(str(output_path)).get_fields()
    assert fields["nome"]["/V"] == "João da Silva"


def test_import_csv_rejects_file_with_no_valid_rows(tmp_path):
    pdf_path = _build_pdf_with_various_fields(tmp_path)
    csv_path = tmp_path / "vazio.csv"
    csv_path.write_text("outra_coluna,valor\n", encoding="utf-8")

    with pytest.raises(ValueError):
        ImportFormDataCSV().run(pdf_path, str(tmp_path / "out.pdf"), str(csv_path))


def test_export_and_import_fdf_round_trip_with_accents(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)
    with_field = tmp_path / "with_field.pdf"
    AddFormField().run(
        str(pdf_path), str(with_field), field_name="nome", page_number=0, x=50, y=300, field_type="text"
    )

    filled = tmp_path / "filled.pdf"
    writer_values_csv = tmp_path / "seed.csv"
    writer_values_csv.write_text("campo,valor\nnome,João Ção\n", encoding="utf-8")
    ImportFormDataCSV().run(str(with_field), str(filled), str(writer_values_csv))

    fdf_path = tmp_path / "dados.fdf"
    count = ExportFormDataFDF().run(str(filled), str(fdf_path))
    assert count == 1
    assert "FEFF" in fdf_path.read_text(encoding="ascii")

    reimported = tmp_path / "reimported.pdf"
    imported_count = ImportFormDataFDF().run(str(with_field), str(reimported), str(fdf_path))
    assert imported_count.filled == 1

    fields = PdfReader(str(reimported)).get_fields()
    assert fields["nome"]["/V"] == "João Ção"


def test_import_fdf_rejects_file_with_no_fields(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1)
    fdf_path = tmp_path / "vazio.fdf"
    fdf_path.write_text("%FDF-1.2\n%%EOF", encoding="ascii")

    with pytest.raises(ValueError):
        ImportFormDataFDF().run(str(pdf_path), str(tmp_path / "out.pdf"), str(fdf_path))


def test_flatten_form_fields_removes_interactivity_but_keeps_value(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)
    with_field = tmp_path / "with_field.pdf"
    AddFormField().run(
        str(pdf_path), str(with_field), field_name="nome", page_number=0, x=50, y=300, width=200, height=20, field_type="text"
    )

    filled = tmp_path / "filled.pdf"
    seed_csv = tmp_path / "seed.csv"
    seed_csv.write_text("campo,valor\nnome,Maria\n", encoding="utf-8")
    ImportFormDataCSV().run(str(with_field), str(filled), str(seed_csv))

    flattened = tmp_path / "flattened.pdf"
    count = FlattenFormFields().run(str(filled), str(flattened))
    assert count == 1

    reader = PdfReader(str(flattened))
    assert not reader.get_fields()
    assert "Maria" in reader.pages[0].extract_text()


def test_flatten_form_fields_rejects_encrypted_input(make_encrypted_pdf, tmp_path):
    encrypted_path = make_encrypted_pdf("protected.pdf", [(200, 200)])

    with pytest.raises(EncryptedPDFError):
        FlattenFormFields().run(str(encrypted_path), str(tmp_path / "out.pdf"))


def _hierarchical_form(tmp_path):
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    pdf_path = tmp_path / "base.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)
    writer = PdfWriter(clone_from=str(pdf_path))
    parent_ref = add_object(writer, DictionaryObject({NameObject("/T"): TextStringObject("cliente")}))
    widget = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/FT"): NameObject("/Tx"),
            NameObject("/T"): TextStringObject("nome"),
            NameObject("/V"): TextStringObject("Maria"),
            NameObject("/Rect"): ArrayObject([FloatObject(v) for v in (100, 300, 300, 320)]),
            NameObject("/Parent"): parent_ref,
            NameObject("/DA"): TextStringObject("/Helv 12 Tf 0 g"),
            NameObject("/P"): writer.pages[0].indirect_reference,
        }
    )
    widget_ref = add_object(writer, widget)
    parent_ref.get_object()[NameObject("/Kids")] = ArrayObject([widget_ref])
    writer.pages[0][NameObject("/Annots")] = ArrayObject([widget_ref])
    helv = DictionaryObject(
        {NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
         NameObject("/BaseFont"): NameObject("/Helvetica")}
    )
    writer.root_object[NameObject("/AcroForm")] = DictionaryObject(
        {
            NameObject("/Fields"): ArrayObject([parent_ref]),
            NameObject("/DR"): DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/Helv"): helv})}),
        }
    )
    path = tmp_path / "hier.pdf"
    with open(path, "wb") as f:
        writer.write(f)
    return path


def test_list_fields_skips_grouping_parents(tmp_path):
    fields = ListFormFields().run(str(_hierarchical_form(tmp_path)))

    assert [(f.name, f.value) for f in fields] == [("cliente.nome", "Maria")]


def test_flatten_hierarchical_fields_removes_the_widget(tmp_path):
    # Em formulários com nomes hierárquicos, o widget continuava interativo por cima do
    # valor estampado.
    flattened = tmp_path / "flat.pdf"
    count = FlattenFormFields().run(str(_hierarchical_form(tmp_path)), str(flattened))

    reader = PdfReader(str(flattened))
    assert count == 1
    assert not reader.get_fields()
    assert "/Annots" not in reader.pages[0]
    assert "Maria" in reader.pages[0].extract_text()


def _form_with_name_field(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)
    with_field = tmp_path / "with_field.pdf"
    AddFormField().run(str(pdf_path), str(with_field), field_name="nome", page_number=0, x=50, y=300)
    return with_field


@pytest.mark.parametrize(
    ("raw", "description"),
    [
        ("campo,valor\nnome,José\n".encode("utf-8-sig"), "CSV UTF-8 com BOM (Excel)"),
        ("campo;valor\nnome;José\n".encode("cp1252"), "CSV do Excel pt-BR (; e Windows-1252)"),
        ("Campo ; Valor\nnome ; José\n".encode("utf-8"), "cabeçalho com maiúsculas e espaços"),
    ],
)
def test_import_csv_accepts_what_excel_saves(tmp_path, raw, description):
    with_field = _form_with_name_field(tmp_path)
    csv_path = tmp_path / "dados.csv"
    csv_path.write_bytes(raw)
    output_path = tmp_path / "out.pdf"

    result = ImportFormDataCSV().run(str(with_field), str(output_path), str(csv_path))

    assert result.filled == 1, description
    assert PdfReader(str(output_path)).get_fields()["nome"]["/V"].strip() == "José"


def test_import_csv_reports_names_that_do_not_exist(tmp_path):
    with_field = _form_with_name_field(tmp_path)
    csv_path = tmp_path / "dados.csv"
    csv_path.write_text("campo,valor\nnome,Ana\nidade,30\n", encoding="utf-8")

    result = ImportFormDataCSV().run(str(with_field), str(tmp_path / "out.pdf"), str(csv_path))

    assert result.filled == 1
    assert result.not_found == ["idade"]


def test_import_csv_fails_when_no_name_matches_instead_of_reporting_success(tmp_path):
    with_field = _form_with_name_field(tmp_path)
    csv_path = tmp_path / "dados.csv"
    csv_path.write_text("campo,valor\nnao_existe,1\n", encoding="utf-8")

    with pytest.raises(ValueError, match="nao_existe"):
        ImportFormDataCSV().run(str(with_field), str(tmp_path / "out.pdf"), str(csv_path))


def test_add_field_rejects_a_name_that_already_exists(tmp_path):
    with_field = _form_with_name_field(tmp_path)

    with pytest.raises(ValueError, match="Já existe um campo"):
        AddFormField().run(
            str(with_field), str(tmp_path / "out.pdf"), field_name="nome", page_number=0, x=50, y=100
        )


def test_add_field_registers_its_font_when_the_form_has_no_default_resources(tmp_path):
    from pypdf.generic import NameObject

    with_field = _form_with_name_field(tmp_path)
    writer = PdfWriter(clone_from=str(with_field))
    del writer.root_object["/AcroForm"][NameObject("/DR")]
    without_dr = tmp_path / "without_dr.pdf"
    with open(without_dr, "wb") as f:
        writer.write(f)

    output_path = tmp_path / "out.pdf"
    AddFormField().run(str(without_dr), str(output_path), field_name="email", page_number=0, x=50, y=100)

    acroform = PdfReader(str(output_path)).trailer["/Root"]["/AcroForm"]
    assert acroform["/DR"]["/Font"]


def test_fdf_export_writes_checkbox_state_as_a_name_and_nests_hierarchical_names(tmp_path):
    form = _build_pdf_with_various_fields(tmp_path)
    filled = tmp_path / "filled.pdf"
    seed = tmp_path / "seed.csv"
    seed.write_text("campo,valor\naceite,/Yes\nnome,Ana\n", encoding="utf-8")
    ImportFormDataCSV().run(str(form), str(filled), str(seed))
    fdf_path = tmp_path / "dados.fdf"

    ExportFormDataFDF().run(str(filled), str(fdf_path))

    text = fdf_path.read_text(encoding="latin-1")
    assert "/V /Yes" in text  # nome, não a string "(/Yes)"
    hierarchical = tmp_path / "hier.fdf"
    ExportFormDataFDF().run(str(_hierarchical_form(tmp_path)), str(hierarchical))
    assert "/Kids [" in hierarchical.read_text(encoding="latin-1")


def test_fdf_import_understands_acrobat_syntax(tmp_path):
    fdf_path = tmp_path / "acrobat.fdf"
    fdf_path.write_bytes(
        b"%FDF-1.2\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /FDF << /Fields [\n"
        # Escapes do próprio FDF: octal (\343 = ã), parênteses escapados e \n.
        rb"<< /T (cliente) /Kids [ << /T (nome) /V (Jo\343o \(filho\)\nSilva) >> ] >>" b"\n"
        b"] >> >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n"
    )
    output = tmp_path / "out.pdf"

    result = ImportFormDataFDF().run(str(_hierarchical_form(tmp_path)), str(output), str(fdf_path))

    assert result.filled == 1
    assert PdfReader(str(output)).get_fields()["cliente.nome"]["/V"] == "João (filho)\nSilva"


def test_fdf_import_reads_checkbox_names(tmp_path):
    form = _build_pdf_with_various_fields(tmp_path)
    fdf_path = tmp_path / "marcado.fdf"
    fdf_path.write_bytes(b"%FDF-1.2\n1 0 obj << /FDF << /Fields [ << /T (aceite) /V /Yes >> ] >> >> endobj\n%%EOF")
    output = tmp_path / "out.pdf"

    ImportFormDataFDF().run(str(form), str(output), str(fdf_path))

    assert PdfReader(str(output)).get_fields()["aceite"]["/V"] == "/Yes"


def test_csv_export_protects_formulas_and_import_round_trips(tmp_path):
    form = _build_pdf_with_various_fields(tmp_path)
    filled = tmp_path / "filled.pdf"
    seed = tmp_path / "seed.csv"
    seed.write_text('campo,valor\nnome,"=HYPERLINK(""http://x"")"\n', encoding="utf-8")
    ImportFormDataCSV().run(str(form), str(filled), str(seed))
    exported = tmp_path / "dados.csv"

    ExportFormDataCSV().run(str(filled), str(exported))

    assert "'=HYPERLINK" in exported.read_text(encoding="utf-8-sig")
    again = tmp_path / "again.pdf"
    ImportFormDataCSV().run(str(form), str(again), str(exported))
    assert PdfReader(str(again)).get_fields()["nome"]["/V"] == '=HYPERLINK("http://x")'


def test_form_data_round_trips_through_xlsx(tmp_path):
    from openpyxl import load_workbook

    form = _build_pdf_with_various_fields(tmp_path)
    filled = tmp_path / "filled.pdf"
    seed = tmp_path / "seed.csv"
    seed.write_text("campo,valor\nnome,=1+1\n", encoding="utf-8")
    ImportFormDataCSV().run(str(form), str(filled), str(seed))
    xlsx = tmp_path / "dados.xlsx"

    ExportFormDataCSV().run(str(filled), str(xlsx))

    cell = next(row for row in load_workbook(xlsx).active.iter_rows() if row[0].value == "nome")[2]
    assert cell.value == "=1+1" and cell.data_type == "s"  # texto, não fórmula
    again = tmp_path / "again.pdf"
    ImportFormDataCSV().run(str(form), str(again), str(xlsx))
    assert PdfReader(str(again)).get_fields()["nome"]["/V"] == "=1+1"


def _tab_order(path) -> list[str]:
    from core.form_fields import _page_widgets, _qualified_name

    return [_qualified_name(widget) for widget in _page_widgets(PdfReader(str(path)).pages[0])]


def test_edit_renames_the_field_and_changes_tooltip_and_flags(tmp_path):
    from core.form_fields import EditFormField

    source = _build_pdf_with_various_fields(tmp_path)
    output = tmp_path / "editado.pdf"

    EditFormField().run(source, str(output), "nome", new_name="nome_completo", tooltip="Como no RG", read_only=True, required=True)

    info = {f.name: f for f in ListFormFields().run(str(output))}
    assert "nome" not in info
    edited = info["nome_completo"]
    assert (edited.field_type, edited.tooltip, edited.read_only, edited.required) == ("text", "Como no RG", True, True)
    assert info["aceite"].read_only is False

    again = tmp_path / "de_novo.pdf"
    EditFormField().run(str(output), str(again), "nome_completo", tooltip="", read_only=False)
    edited = {f.name: f for f in ListFormFields().run(str(again))}["nome_completo"]
    assert (edited.tooltip, edited.read_only, edited.required) == ("", False, True)  # None deixa o obrigatório


def test_edit_renames_only_the_last_part_of_a_hierarchical_name(tmp_path):
    from core.form_fields import EditFormField

    output = tmp_path / "editado.pdf"
    EditFormField().run(str(_hierarchical_form(tmp_path)), str(output), "cliente.nome", new_name="nome_social")

    assert [(f.name, f.value) for f in ListFormFields().run(str(output))] == [("cliente.nome_social", "Maria")]


@pytest.mark.parametrize("new_name, message", [("aceite", "Já existe"), ("a.b", "ponto"), (" ", "vazio")])
def test_edit_refuses_a_bad_new_name(tmp_path, new_name, message):
    from core.form_fields import EditFormField

    with pytest.raises(ValueError, match=message):
        EditFormField().run(_build_pdf_with_various_fields(tmp_path), str(tmp_path / "o.pdf"), "nome", new_name=new_name)


def test_remove_takes_every_widget_of_the_field(tmp_path):
    from core.form_fields import RemoveFormField

    output = tmp_path / "sem_sexo.pdf"
    assert RemoveFormField().run(_build_pdf_with_various_fields(tmp_path), str(output), "sexo") == 2

    assert sorted(PdfReader(str(output)).get_fields()) == ["aceite", "estado", "nome"]
    assert "sexo" not in _tab_order(output)
    with pytest.raises(ValueError, match='Não há campo chamado "sexo"'):
        RemoveFormField().run(str(output), str(tmp_path / "o.pdf"), "sexo")


def test_removing_the_last_field_removes_the_form(tmp_path):
    from core.form_fields import RemoveFormField

    output = tmp_path / "sem_campos.pdf"
    RemoveFormField().run(str(_hierarchical_form(tmp_path)), str(output), "cliente.nome")

    reader = PdfReader(str(output))
    assert "/AcroForm" not in reader.root_object and "/Annots" not in reader.pages[0]


def test_tab_order_by_list_rows_and_columns(tmp_path):
    from core.form_fields import SetTabOrder

    source = _build_pdf_with_various_fields(tmp_path)  # nome (y 350), aceite (300), estado (250), sexo M/F (200)
    by_list, by_rows, by_columns = tmp_path / "lista.pdf", tmp_path / "linhas.pdf", tmp_path / "colunas.pdf"

    SetTabOrder().run(source, str(by_list), order=["estado", "nome"])
    SetTabOrder().run(str(by_list), str(by_rows), mode="rows")
    SetTabOrder().run(source, str(by_columns), mode="columns")

    assert _tab_order(by_list) == ["estado", "nome", "aceite", "sexo", "sexo"]
    assert [f.name for f in ListFormFields().run(str(by_list))] == ["estado", "nome", "aceite", "sexo"]
    assert _tab_order(by_rows) == ["nome", "aceite", "estado", "sexo", "sexo"]
    assert PdfReader(str(by_rows)).pages[0]["/Tabs"] == "/R"
    assert _tab_order(by_columns) == ["nome", "aceite", "estado", "sexo", "sexo"]  # x 50 primeiro, de cima para baixo
    assert PdfReader(str(by_columns)).pages[0]["/Tabs"] == "/C"


def test_date_mask_is_opt_in_javascript(tmp_path):
    pdf_path = tmp_path / "doc.pdf"
    CreateBlankPDF().run(str(pdf_path), page_count=1, width=400, height=400)
    plain, masked = tmp_path / "simples.pdf", tmp_path / "mascara.pdf"
    common = {"page_number": 0, "x": 50, "y": 300, "field_type": "date"}

    AddFormField().run(str(pdf_path), str(plain), field_name="nascimento", **common)
    AddFormField().run(str(pdf_path), str(masked), field_name="nascimento", date_mask=True, **common)

    (plain_widget,) = [a.get_object() for a in PdfReader(str(plain)).pages[0]["/Annots"]]
    (masked_widget,) = [a.get_object() for a in PdfReader(str(masked)).pages[0]["/Annots"]]
    assert "/AA" not in plain_widget
    actions = masked_widget["/AA"]
    assert 'AFDate_KeystrokeEx("dd/mm/yyyy")' in actions["/K"]["/JS"]
    assert 'AFDate_FormatEx("dd/mm/yyyy")' in actions["/F"]["/JS"]


def test_forms_tab_edits_removes_and_orders_the_existing_fields(tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QMessageBox

    from ui.document_session import DocumentSession
    from ui.pages.form_fields_page import FormFieldsPage
    from ui.tasks import wait_for_tasks

    QApplication.instance() or QApplication([])
    monkeypatch.setattr(QMessageBox, "information", lambda *_args: None)
    session = DocumentSession()
    session.open(_build_pdf_with_various_fields(tmp_path))
    page = FormFieldsPage(session)
    outputs = iter(tmp_path / f"passo{n}.pdf" for n in range(1, 4))
    monkeypatch.setattr(page.edit_output_picker, "path", lambda: str(next(outputs)))

    def settle():
        # Cada alteração, ao terminar, lista os campos de novo: uma segunda operação.
        assert any(wait_for_tasks() for _ in range(3))

    def names():
        return [page.fields_list.item(row).data(256) for row in range(page.fields_list.count())]

    page._list_fields()
    settle()
    assert names() == ["nome", "aceite", "estado", "sexo"]

    page.fields_list.setCurrentRow(0)
    assert page.edit_name_edit.text() == "nome"
    page.edit_name_edit.setText("nome_completo")
    page.edit_required_box.setChecked(True)
    page._apply_field_edit()
    settle()
    assert names() == ["nome_completo", "aceite", "estado", "sexo"]  # a lista volta do arquivo novo

    page.fields_list.setCurrentRow(1)
    page._remove_field()
    settle()
    assert names() == ["nome_completo", "estado", "sexo"]

    page.fields_list.insertItem(0, page.fields_list.takeItem(2))  # arrastar "sexo" para o topo
    page._apply_tab_order()
    settle()
    assert names() == ["sexo", "nome_completo", "estado"]
    assert {f.name: f.required for f in ListFormFields().run(session.path())}["nome_completo"] is True


def test_date_mask_option_follows_the_field_type():
    from PySide6.QtWidgets import QApplication

    from ui.document_session import DocumentSession
    from ui.pages.form_fields_page import FormFieldsPage

    QApplication.instance() or QApplication([])
    page = FormFieldsPage(DocumentSession())

    assert not page.date_mask_box.isEnabled() and not page.date_mask_box.isChecked()
    page.type_combo.setCurrentIndex(1)  # Data
    assert page.date_mask_box.isEnabled() and not page.date_mask_box.isChecked()  # desligada por padrão
