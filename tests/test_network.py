"""O app é local (M33): só a verificação de atualização e o carimbo de tempo da assinatura
(opcional, num servidor que o usuário informa) usam a rede.

Duas redes de proteção: uma varredura do código, que acusa qualquer módulo de rede fora
desses dois lugares, e operações que poderiam buscar algo na internet (uma imagem de um HTML,
a cadeia de um certificado) rodadas com a rede bloqueada.
"""

import ast
import os
import socket
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest

ROOT = Path(__file__).resolve().parent.parent

# Módulos (ou o começo deles) que abrem conexões.
_NETWORK_MODULES = (
    "socket", "ssl", "http.client", "urllib.request", "urllib3", "requests", "httpx", "aiohttp",
    "ftplib", "smtplib", "poplib", "imaplib", "telnetlib", "xmlrpc", "PySide6.QtNetwork",
    "PySide6.QtWebEngine", "pyhanko.sign.timestamps",
)  # fmt: skip
# Do Qt Network, só o canal local da instância única (pipe nomeado, socket Unix): não sai da máquina.
_LOCAL_ONLY = {"QLocalServer", "QLocalSocket"}
# Nomes que fazem requisições mesmo vindo de módulos que não são de rede.
_NETWORK_NAMES = {"HTTPTimeStamper", "QNetworkAccessManager", "urlopen"}
# Onde a rede é permitida, e para quê.
_ALLOWED = {
    "ui/main_window.py": "verificação e download de atualização (Ajuda > Verificar atualizações)",
    "core/signature.py": "carimbo de tempo da assinatura, só com o servidor TSA que o usuário informa",
}


def _network_uses(path: Path) -> set[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"))
    found = set()
    for node in ast.walk(tree):
        modules = []
        if isinstance(node, ast.Import):
            modules = [alias.name for alias in node.names]
        elif isinstance(node, ast.ImportFrom) and node.module:
            names = [alias.name for alias in node.names if alias.name not in _LOCAL_ONLY]
            modules = [f"{node.module}.{name}" for name in names]
        found |= {m for m in modules if any(m == n or m.startswith(n + ".") for n in _NETWORK_MODULES)}
        if isinstance(node, ast.Attribute) and node.attr in _NETWORK_NAMES:
            found.add(node.attr)
        if isinstance(node, ast.Name) and node.id in _NETWORK_NAMES:
            found.add(node.id)
    return found


def test_only_the_update_check_and_the_timestamp_use_the_network():
    files = [*ROOT.joinpath("core").rglob("*.py"), *ROOT.joinpath("ui").rglob("*.py"), ROOT / "main.py"]
    offenders = {}
    for path in files:
        relative = path.relative_to(ROOT).as_posix()
        uses = _network_uses(path)
        if uses and relative not in _ALLOWED:
            offenders[relative] = sorted(uses)
    assert offenders == {}
    assert all(_network_uses(ROOT / relative) for relative in _ALLOWED)  # a lista não ficou velha


_LOOPBACK = {"127.0.0.1", "::1", "localhost"}


@pytest.fixture
def no_network(monkeypatch):
    """Toda tentativa de conexão (ou de resolver um nome) fora da máquina fica registrada e
    falha. O próprio computador continua liberado: o asyncio do Windows conversa consigo mesmo
    por um par de sockets locais."""
    attempts = []

    def guard(original, host_of):
        def call(*args, **kwargs):
            host = host_of(args)
            if host in _LOOPBACK:
                return original(*args, **kwargs)
            attempts.append(host)
            raise OSError("rede bloqueada no teste")

        return call

    def address_host(address):
        return address[0] if isinstance(address, tuple) else address

    monkeypatch.setattr(socket.socket, "connect", guard(socket.socket.connect, lambda a: address_host(a[1])))
    monkeypatch.setattr(socket.socket, "connect_ex", guard(socket.socket.connect_ex, lambda a: address_host(a[1])))
    monkeypatch.setattr(socket, "getaddrinfo", guard(socket.getaddrinfo, lambda a: a[0]))
    monkeypatch.setattr(socket, "create_connection", guard(socket.create_connection, lambda a: address_host(a[0])))
    return attempts


def test_signing_and_verifying_stay_offline(no_network, tmp_path):
    from reportlab.pdfgen import canvas as rl_canvas

    from core.signature import DescribeSignatures, SignDocument

    source, signed = tmp_path / "doc.pdf", tmp_path / "assinado.pdf"
    canvas = rl_canvas.Canvas(str(source))
    canvas.drawString(50, 700, "Contrato")
    canvas.save()

    SignDocument().run(str(source), str(signed), test_common_name="Fulano", pades=True)
    (info,) = DescribeSignatures().run(str(signed))  # a cadeia é conferida sem baixar nada

    assert info.valid and no_network == []


def test_html_with_internet_images_and_office_files_stay_offline(no_network, tmp_path):
    from docx import Document
    from PySide6.QtWidgets import QApplication

    from core.office_convert import ConvertOfficeToPDF
    from ui.text_to_pdf import convert_text_to_pdf

    QApplication.instance() or QApplication([])
    html = tmp_path / "pagina.html"
    html.write_text(
        '<h1>Relatório</h1><img src="https://example.com/rastreio.png"><link rel="stylesheet" '
        'href="http://example.com/estilo.css"><p>Texto</p>',
        encoding="utf-8",
    )
    convert_text_to_pdf(str(html), str(tmp_path / "pagina.pdf"))

    docx = tmp_path / "documento.docx"
    document = Document()
    document.add_paragraph("Veja https://example.com")
    document.save(str(docx))
    ConvertOfficeToPDF().run(str(docx), str(tmp_path / "documento.pdf"))

    assert no_network == []


def test_the_guards_do_catch_network_use(no_network, tmp_path):
    # Autoverificação: sem ela, um bloqueio que não bloqueia passaria nos testes acima.
    with pytest.raises(OSError):
        socket.create_connection(("example.com", 80), timeout=1)
    assert no_network == ["example.com"]

    sample = tmp_path / "amostra.py"
    sample.write_text("import urllib.request\nfrom PySide6.QtNetwork import QLocalSocket, QNetworkRequest\n")
    assert _network_uses(sample) == {"urllib.request", "PySide6.QtNetwork.QNetworkRequest"}
