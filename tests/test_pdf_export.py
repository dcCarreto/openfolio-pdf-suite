"""Mais conversões (A09): PDF para Word, Excel e texto com layout; HTML, Markdown e texto
para PDF."""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pypdfium2 as pdfium
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter
from PySide6.QtWidgets import QApplication, QMessageBox, QPushButton
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table

from core.extract_text import ExtractText
from core.fonts import SANS, register_fonts
from core.pdf_export import ConvertPDFToExcel, ConvertPDFToWord, cell_value, paragraphs, table_rows
from ui.document_session import DocumentSession
from ui.pages.convert_page import ConvertPage
from ui.tasks import wait_for_tasks
from ui.text_to_pdf import convert_text_to_pdf

_TEXT = (
    "Este relatório resume as vendas do trimestre, com os valores por região e o total geral. "
    "Os números são preliminares e podem mudar depois da auditoria, que acontece sempre no mês "
    "seguinte ao fechamento."
)
_TABLE = [
    ["Região", "Janeiro", "Fevereiro", "Total"],
    ["Norte", "1.234,56", "980", "2.214,56"],
    ["Sul", "12,5", "", "12,5"],
    ["Centro-Oeste", "=SOMA(A1)", "77", "77"],
]


@pytest.fixture
def report(tmp_path):
    """Página 1: título, parágrafo e tabela; página 2: em branco."""
    register_fonts()
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontName = SANS
    path = tmp_path / "relatorio.pdf"
    SimpleDocTemplate(str(path), pagesize=A4).build(
        [
            Paragraph("Relatório de vendas", styles["Title"]),
            Paragraph(_TEXT, body),
            Paragraph("Segundo parágrafo, curto.", body),
            Table(_TABLE, style=[("FONTNAME", (0, 0), (-1, -1), SANS)]),
        ]
    )
    writer = PdfWriter(clone_from=str(path))
    writer.add_blank_page()
    writer.write(path)
    return path


def test_paragraphs_rejoin_lines_and_hyphenated_words():
    text = "Título\numa linha longa que continua na\nseguinte e termina aqui com uma pala-\nvra quebrada no fim.\n\nOutro."
    assert paragraphs(text) == [
        "Título",
        "uma linha longa que continua na seguinte e termina aqui com uma palavra quebrada no fim.",
        "Outro.",
    ]


def test_table_rows_keep_empty_cells_in_place():
    layout = "  Texto corrido\n  Nome     Jan     Fev\n  Sul      12,5            \n  Leste            3\n"
    assert table_rows(layout) == [["Texto corrido"], ["Nome", "Jan", "Fev"], ["Sul", "12,5"], ["Leste", "", "3"]]


@pytest.mark.parametrize(
    ("text", "value"),
    [("980", 980), ("-7", -7), ("1.234,56", 1234.56), ("12,5", 12.5), ("12.5", 12.5),
     ("1.234", "1.234"), ("007", "007"), ("12345678901234567", "12345678901234567"), ("=SOMA(A1)", "=SOMA(A1)")],
)
def test_cell_value_converts_only_unambiguous_numbers(text, value):
    assert cell_value(text) == value


def test_pdf_to_word_writes_paragraphs_and_reports_pages_without_text(report, tmp_path):
    output = tmp_path / "relatorio.docx"

    assert ConvertPDFToWord().run(str(report), str(output)) == [1]

    texts = [paragraph.text for paragraph in Document(output).paragraphs if paragraph.text]
    assert texts[:3] == ["Relatório de vendas", _TEXT, "Segundo parágrafo, curto."]


def test_pdf_to_excel_puts_cells_in_columns_with_numbers_and_no_formulas(report, tmp_path):
    output = tmp_path / "relatorio.xlsx"

    assert ConvertPDFToExcel().run(str(report), str(output)) == [1]

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["1", "2"]
    rows = [list(row) for row in workbook["1"].iter_rows(values_only=True)]
    assert rows[0][0] == "Relatório de vendas"
    table = rows[-4:]
    assert table[0][:4] == ["Região", "Janeiro", "Fevereiro", "Total"]
    assert table[1][:4] == ["Norte", 1234.56, 980, 2214.56]
    assert table[2][:4] == ["Sul", 12.5, None, 12.5]
    assert table[3][1] == "=SOMA(A1)" and workbook["1"].cell(row=len(rows), column=2).data_type == "s"


def test_control_characters_in_the_pdf_text_do_not_break_word_or_excel(report, tmp_path, monkeypatch):
    # O PDF que o Qt gera (e outros) traz caracteres de controle que o XML do Office recusa.
    monkeypatch.setattr("core.pdf_export.page_text", lambda _page, layout=False: "Texto\x01 com\tcontrole\x0b\n")

    ConvertPDFToWord().run(str(report), str(tmp_path / "a.docx"))
    ConvertPDFToExcel().run(str(report), str(tmp_path / "a.xlsx"))

    assert Document(tmp_path / "a.docx").paragraphs[0].text == "Texto com controle"
    assert load_workbook(tmp_path / "a.xlsx")["1"]["A1"].value == "Texto com controle"


def test_text_with_layout_keeps_columns_aligned(report, tmp_path):
    output = tmp_path / "relatorio.txt"
    ExtractText().run(str(report), str(output), layout=True)
    lines = output.read_text(encoding="utf-8").splitlines()
    header = next(line for line in lines if "Janeiro" in line)
    south = next(line for line in lines if line.strip().startswith("Sul"))
    assert south.index("12,5") == header.index("Janeiro")
    assert south.rindex("12,5") == header.index("Total")


def _app():
    return QApplication.instance() or QApplication([])


def _pdf_text(path):
    # Pelo pdfium, como a busca do app: o pypdf lê o espaço entre palavras do Qt como tabulação.
    document = pdfium.PdfDocument(str(path))
    try:
        return "\n".join(page.get_textpage().get_text_range() for page in document)
    finally:
        document.close()


def test_html_markdown_and_text_become_searchable_pdfs(tmp_path):
    _app()
    Image.new("RGB", (40, 20), "red").save(tmp_path / "logo.png")
    html = tmp_path / "pagina.html"
    html.write_text(
        "<html><head><title>Página</title></head><body><h1>Proposta comercial</h1>"
        "<p>Valor <b>total</b>: R$ 1.000</p><img src='logo.png'>"
        "<img src='https://exemplo.invalid/rastreador.png'><script>alert(1)</script>"
        "<table><tr><td>Item</td><td>Preço</td></tr></table></body></html>",
        encoding="utf-8",
    )
    markdown = tmp_path / "notas.md"
    markdown.write_text("# Ata da reunião\n\n- primeiro item\n- segundo item\n\n| A | B |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8")
    text = tmp_path / "antigo.txt"
    text.write_bytes("Linha em latin-1: ação\n".encode("latin-1") + "texto longo ".encode() * 400)

    for source in (html, markdown, text):
        convert_text_to_pdf(str(source), str(tmp_path / f"{source.stem}.pdf"))

    html_text = _pdf_text(tmp_path / "pagina.pdf")
    assert "Proposta comercial" in html_text and "Preço" in html_text and "alert" not in html_text
    assert PdfReader(tmp_path / "pagina.pdf").metadata.title == "Página"
    # A imagem local entrou de verdade (relativa ao HTML, não à pasta atual), não o ícone de
    # imagem quebrada; a remota não foi buscada.
    images = [image.image for image in PdfReader(tmp_path / "pagina.pdf").pages[0].images]
    assert [image.size for image in images] == [(40, 20)]
    red, green, blue = images[0].convert("RGB").getpixel((5, 5))
    assert red > 240 and green < 15 and blue < 15  # o Qt recomprime: 254 em vez de 255
    markdown_text = _pdf_text(tmp_path / "notas.pdf")
    assert "Ata da reunião" in markdown_text and "segundo item" in markdown_text and "#" not in markdown_text
    assert "ação" in _pdf_text(tmp_path / "antigo.pdf")
    assert len(PdfReader(tmp_path / "antigo.pdf").pages) > 1


def test_markdown_with_unclosed_html_tags_keeps_the_following_text(tmp_path):
    _app()
    Image.new("RGB", (40, 20), "red").save(tmp_path / "icone.png")
    source = tmp_path / "LEIAME.md"
    source.write_text(
        '<p align="center">\n  <img src="icone.png" width="40">\n</p>\n\n'
        '<p align="center">\n  primeira linha<br>\n  segunda linha\n</p>\n\n'
        "## Funcionalidades\n\n| Aba | O que faz |\n|---|---|\n| [Mesclar](https://exemplo.invalid) | Junta **PDFs** |\n",
        encoding="utf-8",
    )
    convert_text_to_pdf(str(source), str(tmp_path / "LEIAME.pdf"))

    text = _pdf_text(tmp_path / "LEIAME.pdf")
    assert "segunda linha" in text and "Funcionalidades" in text and "Junta PDFs" in text and "<br" not in text


def test_convert_page_exports_the_open_document(report, tmp_path, monkeypatch):
    _app()
    messages = []
    monkeypatch.setattr(QMessageBox, "information", lambda _parent, _title, text: messages.append(text))
    session = DocumentSession()
    session.open(str(report))
    page = ConvertPage(session)

    page.direction_combo.setCurrentIndex(page.direction_combo.findText("PDF para Excel (.xlsx)"))
    output = page.excel_output_picker.path()
    assert output == str(tmp_path / "relatorio_convertido.xlsx")
    next(button for button in page.stack.currentWidget().findChildren(QPushButton) if button.text() == "Converter").click()
    wait_for_tasks()

    assert os.path.exists(output)
    assert "Páginas sem texto (escaneadas? faça o OCR antes): 2" in messages[-1]
